"""
DOCX export service using python-docx.
Generates ATS-safe .docx from structured resume data.
"""
import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from app.schemas import ResumeOut


def add_horizontal_rule(doc: Document):
    """Add a simple bottom border to simulate a horizontal rule."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_heading(doc: Document, text: str):
    """Add an ATS-safe section heading (no tables, no text boxes)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Arial"
    # Add bottom border to the paragraph
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_docx(resume: ResumeOut) -> bytes:
    """Build an ATS-safe DOCX from structured resume data."""
    doc = Document()

    # Page margins
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    # Default paragraph font
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    contact = resume.contact
    name = contact.full_name if contact else "Your Name"

    # Name (centered, bold, large)
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_para.paragraph_format.space_after = Pt(2)
    name_run = name_para.add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_run.font.name = "Arial"

    # Contact line (centered)
    if contact:
        contact_parts = []
        for field in [contact.email, contact.phone, contact.city, contact.linkedin, contact.github, contact.portfolio]:
            if field:
                contact_parts.append(field)
        if contact_parts:
            contact_para = doc.add_paragraph(" | ".join(contact_parts))
            contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact_para.paragraph_format.space_before = Pt(0)
            contact_para.paragraph_format.space_after = Pt(4)
            for run in contact_para.runs:
                run.font.size = Pt(9.5)
                run.font.name = "Arial"

    # Sections
    for section_obj in resume.sections:
        section_type = section_obj.type.value
        heading_map = {
            "education": "Education",
            "experience": "Experience",
            "project": "Projects",
            "skill": "Skills",
            "achievement": "Achievements",
        }
        heading = heading_map.get(section_type, section_type.title())
        add_section_heading(doc, heading)

        for entry in section_obj.entries:
            p = entry.payload

            if section_type == "education":
                # Institution + dates on same line
                row_p = doc.add_paragraph()
                row_p.paragraph_format.space_before = Pt(4)
                row_p.paragraph_format.space_after = Pt(0)
                run_inst = row_p.add_run(p.get("institution", ""))
                run_inst.bold = True
                run_inst.font.size = Pt(10.5)
                date_str = f"{p.get('start_date','')} - {p.get('end_date','')}"
                run_date = row_p.add_run(f"\t{date_str}")
                run_date.font.size = Pt(10)
                row_p.paragraph_format.tab_stops.add_tab_stop(Inches(5.5))

                degree_p = doc.add_paragraph(f"{p.get('degree','')} in {p.get('field','')}")
                degree_p.runs[0].italic = True
                degree_p.runs[0].font.size = Pt(10)
                degree_p.paragraph_format.space_before = Pt(0)
                if p.get("gpa"):
                    gpa_p = doc.add_paragraph(f"GPA: {p['gpa']}")
                    gpa_p.runs[0].font.size = Pt(10)

            elif section_type in ("experience", "project"):
                title = p.get("company", p.get("name", ""))
                subtitle = p.get("role", p.get("tech_stack", ""))
                date_str = f"{p.get('start_date','')} - {p.get('end_date','Present')}"
                location = p.get("location", "")

                row_p = doc.add_paragraph()
                row_p.paragraph_format.space_before = Pt(4)
                row_p.paragraph_format.space_after = Pt(0)
                run_title = row_p.add_run(title)
                run_title.bold = True
                run_title.font.size = Pt(10.5)
                run_date = row_p.add_run(f"\t{date_str}")
                run_date.font.size = Pt(10)
                row_p.paragraph_format.tab_stops.add_tab_stop(Inches(5.5))

                sub_text = subtitle
                if location:
                    sub_text += f", {location}"
                if sub_text:
                    sub_p = doc.add_paragraph(sub_text)
                    sub_p.runs[0].italic = True
                    sub_p.runs[0].font.size = Pt(10)
                    sub_p.paragraph_format.space_before = Pt(0)

                for bullet in entry.bullets:
                    text = bullet.improved_text or bullet.raw_text
                    b_para = doc.add_paragraph(text, style="List Bullet")
                    b_para.paragraph_format.space_before = Pt(1)
                    b_para.paragraph_format.space_after = Pt(1)
                    b_para.paragraph_format.left_indent = Inches(0.2)
                    for run in b_para.runs:
                        run.font.size = Pt(10.5)
                        run.font.name = "Arial"

            elif section_type == "skill":
                groups = p.get("groups", {})
                if isinstance(groups, dict):
                    for group_name, skills in groups.items():
                        sk_p = doc.add_paragraph()
                        sk_p.paragraph_format.space_before = Pt(2)
                        sk_p.paragraph_format.space_after = Pt(1)
                        run_cat = sk_p.add_run(f"{group_name}: ")
                        run_cat.bold = True
                        run_cat.font.size = Pt(10.5)
                        run_skills = sk_p.add_run(str(skills))
                        run_skills.font.size = Pt(10.5)

            elif section_type == "achievement":
                ach_p = doc.add_paragraph()
                ach_p.paragraph_format.space_before = Pt(4)
                run_title = ach_p.add_run(p.get("title", ""))
                run_title.bold = True
                run_title.font.size = Pt(10.5)
                date_str = p.get("date", "")
                if date_str:
                    ach_p.add_run(f"\t{date_str}").font.size = Pt(10)
                if p.get("issuer"):
                    iss_p = doc.add_paragraph(p["issuer"])
                    iss_p.runs[0].italic = True
                    iss_p.runs[0].font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
